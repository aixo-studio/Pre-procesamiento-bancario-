"""Prueba de la entrada: convertir un archivo con su banco, verificar que sea el
documento indicado, guardar en la carpeta del banco, borrar resultados viejos y
avisar fechas de otro mes. Todo con archivos inventados.
"""
from datetime import date, datetime

import openpyxl
import pytest

from convertir import convertir_documento
from motor.armar import Movimiento
from motor.carpetas import fechas_fuera_del_mes, ruta_salida

ENCABEZADO_BBVA = ["Fecha", "Fecha Valor", "Concepto", "Codigo", "Número Documento",
                   "Oficina", "Crédito", "Débito", "Detalle"]
ENCABEZADO_GALICIA = ["Fecha", "Descripción", "Origen", "Débitos", "Créditos", "Grupo de Conceptos",
                      "Concepto", "Número de Terminal", "Observaciones Cliente", "Número de Comprobante",
                      "Leyendas Adicionales 1", "Leyendas Adicionales 2", "Leyendas Adicionales 3",
                      "Leyendas Adicionales 4", "Tipo de Movimiento", "Saldo"]
ENCABEZADO_ERP = ["Año - mes", "Codigo cuenta", "Cuenta", "Fecha", "Comprobante", "Descripción",
                  "Debe", "Haber", "Documento", "Dimensión valor", "Detalle", "Cotizacion",
                  "Productoid", "Producto", "Empresa", "Organizacion transaccion",
                  "Debe", "Haber", "Tipo de documento"]


def guardar(ruta, filas):
    """Guarda las filas en un Excel nuevo y devuelve la ruta."""
    libro = openpyxl.Workbook()
    for fila in filas:
        libro.active.append(fila)
    libro.save(ruta)
    return ruta


def extracto_bbva(ruta, fecha="30-09-2026"):
    """Un extracto del BBVA inventado, con una fila."""
    return guardar(ruta, [["Empresa:", "EMPRESA DE PRUEBA"], [], ENCABEZADO_BBVA,
                          [fecha, fecha, "TRANSFERENCI", "319", "", "OF", "", -1500, ""]])


def libro_erp(ruta):
    """Un libro del ERP inventado: saldo inicial y un movimiento."""
    fila = ["2026-09", "1", "BANCO", datetime(2026, 9, 1), "F", "", 0, 0, "Saldo Inicial",
            None, "", 1, None, None, "EMPRESA", "", 0, 0, ""]
    movimiento = fila[:6] + [0, 1500, "PAGO - 1"] + fila[9:16] + [0, 1, "Orden de pago"]
    return guardar(ruta, [ENCABEZADO_ERP, fila, movimiento])


def test_convierte_extracto_y_libro_a_la_carpeta_del_banco(tmp_path):
    output = tmp_path / "output"
    resumen = convertir_documento(extracto_bbva(tmp_path / "extracto.xlsx"), "BBVA", False, "2026-09", output)
    assert resumen["filas"] == 1 and resumen["suma"] == -150000
    resumen = convertir_documento(libro_erp(tmp_path / "libro.xlsx"), "BBVA", True, "2026-09", output)
    assert resumen["filas"] == 1 and resumen["ignoradas"] == 1
    salida = output / "2026-09" / "BBVA"
    assert sorted(ruta.name for ruta in salida.iterdir()) == \
        ["banco_bbva - extracto.xlsx", "sistema_finnegans - libro.xlsx"]
    libro = openpyxl.load_workbook(salida / "sistema_finnegans - libro.xlsx")
    assert libro.sheetnames == ["Movimientos"]
    assert list(libro.active.iter_rows(values_only=True))[1][1:] == ("Orden de pago", None, -1500)


def test_frena_si_el_banco_no_esta_configurado(tmp_path):
    with pytest.raises(ValueError, match='"Macro" no está en la configuración'):
        convertir_documento(extracto_bbva(tmp_path / "a.xlsx"), "Macro", False, "2026-09", tmp_path)


def test_frena_si_el_extracto_es_de_otro_banco(tmp_path):
    ruta = guardar(tmp_path / "galicia.xlsx", [ENCABEZADO_GALICIA])
    with pytest.raises(ValueError, match="no es un extracto de BBVA"):
        convertir_documento(ruta, "BBVA", False, "2026-09", tmp_path)


def test_frena_si_se_confunde_extracto_y_libro(tmp_path):
    with pytest.raises(ValueError, match="no es el libro del ERP"):
        convertir_documento(extracto_bbva(tmp_path / "a.xlsx"), "BBVA", True, "2026-09", tmp_path)
    with pytest.raises(ValueError, match="no es un extracto de BBVA"):
        convertir_documento(libro_erp(tmp_path / "b.xlsx"), "BBVA", False, "2026-09", tmp_path)


def test_frena_si_el_archivo_no_se_puede_abrir(tmp_path):
    ruta = tmp_path / "roto.xlsx"
    ruta.write_text("no es un Excel")
    with pytest.raises(ValueError, match="no se pudo abrir como Excel"):
        convertir_documento(ruta, "BBVA", False, "2026-09", tmp_path)


def test_frena_si_el_mes_no_tiene_la_forma_correcta(tmp_path):
    with pytest.raises(ValueError, match="AAAA-MM"):
        convertir_documento(extracto_bbva(tmp_path / "a.xlsx"), "BBVA", False, "09-2026", tmp_path)


def test_si_falla_borra_el_resultado_viejo(tmp_path):
    output = tmp_path / "output"
    ruta = extracto_bbva(tmp_path / "extracto.xlsx")
    convertir_documento(ruta, "BBVA", False, "2026-09", output)
    # El extracto ahora trae una fecha que no se puede leer.
    extracto_bbva(ruta, fecha="sin fecha")
    with pytest.raises(ValueError, match="Se borró el resultado anterior: banco_bbva - extracto.xlsx"):
        convertir_documento(ruta, "BBVA", False, "2026-09", output)
    assert list((output / "2026-09" / "BBVA").iterdir()) == []


def test_avisa_fechas_de_otro_mes_sin_sacarlas(tmp_path):
    output = tmp_path / "output"
    resumen = convertir_documento(extracto_bbva(tmp_path / "a.xlsx", fecha="31-08-2026"),
                                  "BBVA", False, "2026-09", output)
    assert resumen["fuera_del_mes"] == {"2026-08": 1}
    assert openpyxl.load_workbook(resumen["ruta_salida"]).active.max_row == 2


def test_nombre_de_salida_sin_la_extension_del_original(tmp_path):
    assert ruta_salida(tmp_path, "banco_bbva", tmp_path / "banco X.xls").name == "banco_bbva - banco X.xlsx"


def test_cuenta_fechas_fuera_del_mes():
    movimientos = [Movimiento(date(2026, 9, 1), "", "", 1), Movimiento(date(2026, 8, 31), "", "", 1),
                   Movimiento(date(2026, 10, 1), "", "", 1), Movimiento(date(2026, 8, 30), "", "", 1)]
    assert fechas_fuera_del_mes(movimientos, "2026-09") == {"2026-08": 2, "2026-10": 1}
