"""Prueba del extracto del BBVA con un archivo inventado que imita la forma del real."""
from datetime import date

import openpyxl
import pytest

from motor.armar import armar_movimientos
from motor.configuracion import buscar_reglas_y_carpeta
from motor.escribir import controlar, escribir_excel
from motor.leer import encontrar_tabla, leer_filas

reglas, _ = buscar_reglas_y_carpeta("BBVA", es_erp=False)

ENCABEZADO = ["Fecha", "Fecha Valor", "Concepto", "Codigo", "Número Documento",
              "Oficina", "Crédito", "Débito", "Detalle", "   "]


def crear_extracto(carpeta, filas_de_datos):
    """Arma un Excel con títulos arriba, el encabezado y los movimientos."""
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.append(["Empresa: ", "EMPRESA DE PRUEBA"])
    hoja.append(["Cuenta: ", "000-000000/0"])
    hoja.append([])
    hoja.append(ENCABEZADO)
    for fila in filas_de_datos:
        hoja.append(fila)
    ruta = carpeta / "extracto.xlsx"
    libro.save(ruta)
    return ruta


def convertir(ruta):
    """Corre los pasos del motor sobre el archivo de prueba y devuelve movimientos y resumen."""
    columnas, datos = encontrar_tabla(leer_filas(ruta), reglas["encabezados"])
    movimientos = armar_movimientos(columnas, datos, reglas)
    return movimientos, controlar(movimientos, columnas, datos, reglas)


def test_convierte_a_cuatro_columnas(tmp_path):
    ruta = crear_extracto(tmp_path, [
        ["30-09-2026", "30-09-2026", "CH/CLEAR.48  000000000001", "005", "90000001", "OF", "", -150000, "", ""],
        ["29-09-2026", "29-09-2026", "TRANSFERENCI 20111111112", "319", "", "OF", 250000, "", "CTE 000011110001", ""],
        [],  # una fila vacía en el medio se saltea
        ["01-09-2026", "01-09-2026", "IVA TASA GRA", "236", "", "OF", "", -1365.5, "", ""],
    ])
    movimientos, resumen = convertir(ruta)

    assert [m.fecha for m in movimientos] == [date(2026, 9, 30), date(2026, 9, 29), date(2026, 9, 1)]
    # El Código sale de la columna Codigo, como texto y con sus ceros adelante.
    assert [m.codigo for m in movimientos] == ["005", "319", "236"]
    assert [m.descripcion for m in movimientos] == ["", "CTE 000011110001", ""]
    assert [m.monto_centavos for m in movimientos] == [-15000000, 25000000, -136550]
    assert (resumen["filas"], resumen["ignoradas"], resumen["creditos"], resumen["debitos"], resumen["suma"]) == \
        (3, 0, 25000000, 15136550, 9863450)

    salida = tmp_path / "salida.xlsx"
    escribir_excel(movimientos, salida)
    libro = openpyxl.load_workbook(salida)
    assert libro.sheetnames == ["Movimientos"]  # una sola hoja, con todo
    filas = list(libro.active.iter_rows(values_only=True))
    assert filas[0] == ("Fecha", "Código", "Descripción", "Monto")
    assert filas[2][1:] == ("319", "CTE 000011110001", 250000)
    # En el Excel el código con ceros sigue siendo el texto "005", no el número 5.
    assert filas[1][1] == "005" and isinstance(filas[1][1], str)


def test_frena_si_la_fecha_no_se_puede_leer(tmp_path):
    ruta = crear_extracto(tmp_path, [
        ["2026/09/30", "", "TRANSFERENCI", "319", "", "OF", 100, "", "", ""],
    ])
    with pytest.raises(ValueError, match="Fila 5: fecha"):
        convertir(ruta)


def test_frena_si_una_fila_tiene_credito_y_debito(tmp_path):
    ruta = crear_extracto(tmp_path, [
        ["30-09-2026", "", "TRANSFERENCI", "319", "", "OF", 100, -100, "", ""],
    ])
    with pytest.raises(ValueError, match="Fila 5: tiene que tener crédito o débito"):
        convertir(ruta)


def test_frena_si_no_es_un_extracto_del_bbva(tmp_path):
    libro = openpyxl.Workbook()
    libro.active.append(["Fecha", "Descripción", "Importe"])
    ruta = tmp_path / "otro.xlsx"
    libro.save(ruta)
    with pytest.raises(ValueError, match="No se encontró la fila de encabezados"):
        convertir(ruta)


def test_frena_si_el_codigo_viene_como_numero(tmp_path):
    # Si alguien abrió y guardó el archivo en Excel, "005" puede haber quedado como el número 5.
    ruta = crear_extracto(tmp_path, [
        ["30-09-2026", "30-09-2026", "CH/CLEAR.48", 5, "", "OF", "", -150000, "", ""],
    ])
    with pytest.raises(ValueError, match="Fila 5: el Código viene como número"):
        convertir(ruta)
