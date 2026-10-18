"""Prueba de las carpetas: reconocer cada archivo, revisar la carpeta del banco,
nombrar y borrar resultados, y avisar fechas de otro mes. Todo con archivos inventados.
"""
from datetime import date, datetime

import openpyxl
import pytest

from convertir import convertir_banco
from motor.armar import Movimiento
from motor.carpetas import fechas_fuera_del_mes, reconocer_documento, revisar_carpeta, ruta_salida
from motor.leer import leer_filas

ENCABEZADO_BBVA = ["Fecha", "Fecha Valor", "Concepto", "Codigo", "Número Documento",
                   "Oficina", "Crédito", "Débito", "Detalle"]
ENCABEZADO_GALICIA = ["Fecha", "Descripción", "Origen", "Débitos", "Créditos", "Grupo de Conceptos",
                      "Concepto", "Número de Terminal", "Observaciones Cliente", "Número de Comprobante",
                      "Leyendas Adicionales 1", "Leyendas Adicionales 2", "Leyendas Adicionales 3",
                      "Leyendas Adicionales 4", "Tipo de Movimiento", "Saldo"]
ENCABEZADO_SISTEMA = ["Año - mes", "Codigo cuenta", "Cuenta", "Fecha", "Comprobante", "Descripción",
                      "Debe", "Haber", "Documento", "Dimensión valor", "Detalle", "Cotizacion",
                      "Productoid", "Producto", "Empresa", "Organizacion transaccion",
                      "Debe", "Haber", "Tipo de documento"]


def guardar(ruta, filas):
    """Guarda las filas en un Excel nuevo y devuelve la ruta."""
    libro = openpyxl.Workbook()
    for fila in filas:
        libro.active.append(fila)
    libro.save(ruta)
    return ruta


def extracto_bbva(ruta, fecha="30-09-2026", debito=-1500):
    """Un extracto del BBVA inventado, con una fila."""
    return guardar(ruta, [["Empresa:", "EMPRESA DE PRUEBA"], [], ENCABEZADO_BBVA,
                          [fecha, fecha, "TRANSFERENCI", "319", "", "OF", "", debito, ""]])


def libro_sistema(ruta):
    """Un libro del sistema inventado: saldo inicial y un movimiento."""
    fila = ["2026-09", "1", "BANCO", datetime(2026, 9, 1), "F", "", 0, 0, "Saldo Inicial",
            None, "", 1, None, None, "EMPRESA", "", 0, 0, ""]
    movimiento = fila[:6] + [0, 1500, "PAGO - 1"] + fila[9:16] + [0, 1, "Orden de pago"]
    return guardar(ruta, [ENCABEZADO_SISTEMA, fila, movimiento])


def carpeta_bbva(base):
    """Arma input/BBVA con un extracto y un libro del sistema."""
    carpeta = base / "input" / "BBVA"
    carpeta.mkdir(parents=True)
    extracto_bbva(carpeta / "extracto.xlsx")
    libro_sistema(carpeta / "libro.xlsx")
    return carpeta


def test_reconoce_cada_documento_por_sus_encabezados(tmp_path):
    assert reconocer_documento(leer_filas(extracto_bbva(tmp_path / "a.xlsx")), "BBVA") == "banco_bbva"
    assert reconocer_documento(leer_filas(libro_sistema(tmp_path / "b.xlsx")), "BBVA") == "sistema_finnegans"


def test_frena_si_el_extracto_es_de_otro_banco(tmp_path):
    filas = leer_filas(guardar(tmp_path / "galicia.xlsx", [ENCABEZADO_GALICIA]))
    with pytest.raises(ValueError, match='extracto de banco_galicia, pero está en la carpeta "BBVA"'):
        reconocer_documento(filas, "BBVA")


def test_frena_si_no_reconoce_el_archivo(tmp_path):
    filas = leer_filas(guardar(tmp_path / "otro.xlsx", [["Fecha", "Importe"]]))
    with pytest.raises(ValueError, match="no se reconoce"):
        reconocer_documento(filas, "BBVA")


def test_revisa_que_haya_un_extracto_y_un_libro(tmp_path):
    carpeta = carpeta_bbva(tmp_path)
    (carpeta / ".DS_Store").write_text("")  # los archivos ocultos no cuentan
    assert sorted(documento for _, documento, _ in revisar_carpeta(carpeta)) == ["banco_bbva", "sistema_finnegans"]
    # Uno de más
    extracto_bbva(carpeta / "extracto repetido.xlsx")
    with pytest.raises(ValueError, match="hay 2 y 1"):
        revisar_carpeta(carpeta)
    # Uno que falta
    (carpeta / "extracto repetido.xlsx").unlink()
    (carpeta / "libro.xlsx").unlink()
    with pytest.raises(ValueError, match="hay 1 y 0"):
        revisar_carpeta(carpeta)


def test_frena_si_el_banco_no_es_conocido(tmp_path):
    carpeta = tmp_path / "Macro"
    carpeta.mkdir()
    with pytest.raises(ValueError, match="no es un banco conocido"):
        revisar_carpeta(carpeta)


def test_nombre_de_salida_sin_la_extension_del_original(tmp_path):
    assert ruta_salida(tmp_path, "banco_bbva", tmp_path / "banco X.xls").name == "banco_bbva - banco X.xlsx"


def test_convierte_cada_archivo_por_separado(tmp_path):
    carpeta = carpeta_bbva(tmp_path)
    salida = tmp_path / "output" / "BBVA"
    assert convertir_banco(carpeta, salida, "2026-09") == (2, 0)
    assert sorted(ruta.name for ruta in salida.iterdir()) == \
        ["banco_bbva - extracto.xlsx", "sistema_finnegans - libro.xlsx"]
    libro = openpyxl.load_workbook(salida / "sistema_finnegans - libro.xlsx")
    assert libro.sheetnames == ["Movimientos"]
    assert list(libro.active.iter_rows(values_only=True))[1][1:] == ("Orden de pago", None, -1500)


def test_si_un_archivo_falla_borra_su_resultado_viejo_y_los_demas_siguen(tmp_path):
    carpeta = carpeta_bbva(tmp_path)
    salida = tmp_path / "output" / "BBVA"
    convertir_banco(carpeta, salida, "2026-09")
    # El extracto ahora trae una fecha que no se puede leer.
    extracto_bbva(carpeta / "extracto.xlsx", fecha="sin fecha")
    assert convertir_banco(carpeta, salida, "2026-09") == (1, 1)
    assert [ruta.name for ruta in salida.iterdir()] == ["sistema_finnegans - libro.xlsx"]


def test_si_la_carpeta_esta_mal_borra_todos_sus_resultados(tmp_path):
    carpeta = carpeta_bbva(tmp_path)
    salida = tmp_path / "output" / "BBVA"
    convertir_banco(carpeta, salida, "2026-09")
    (carpeta / "libro.xlsx").unlink()
    assert convertir_banco(carpeta, salida, "2026-09") == (0, 1)
    assert list(salida.iterdir()) == []


def test_un_solo_archivo(tmp_path):
    carpeta = carpeta_bbva(tmp_path)
    salida = tmp_path / "output" / "BBVA"
    assert convertir_banco(carpeta, salida, "2026-09", "libro.xlsx") == (1, 0)
    assert [ruta.name for ruta in salida.iterdir()] == ["sistema_finnegans - libro.xlsx"]


def test_avisa_fechas_de_otro_mes_sin_sacarlas(tmp_path, capsys):
    carpeta = carpeta_bbva(tmp_path)
    extracto_bbva(carpeta / "extracto.xlsx", fecha="31-08-2026")
    salida = tmp_path / "output" / "BBVA"
    assert convertir_banco(carpeta, salida, "2026-09") == (2, 0)
    assert "⚠ 1 fila con fecha de 2026-08" in capsys.readouterr().out
    assert openpyxl.load_workbook(salida / "banco_bbva - extracto.xlsx").active.max_row == 2


def test_cuenta_fechas_fuera_del_mes():
    movimientos = [Movimiento(date(2026, 9, 1), "", "", 1), Movimiento(date(2026, 8, 31), "", "", 1),
                   Movimiento(date(2026, 10, 1), "", "", 1), Movimiento(date(2026, 8, 30), "", "", 1)]
    assert fechas_fuera_del_mes(movimientos, "2026-09") == {"2026-08": 2, "2026-10": 1}
