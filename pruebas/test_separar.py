"""Prueba de la separación de gastos bancarios: cada fila va a una sola hoja,
en el orden del original; el extracto de un banco sale siempre con dos hojas
y el libro del ERP con una. Todo con datos inventados.
"""
from datetime import date, datetime

import openpyxl

from convertir import convertir_documento
from motor.armar import Movimiento
from motor.separar import separar_gastos

ENCABEZADO_BBVA = ["Fecha", "Fecha Valor", "Concepto", "Codigo", "Número Documento",
                   "Oficina", "Crédito", "Débito", "Detalle"]
ENCABEZADO_ERP = ["Año - mes", "Codigo cuenta", "Cuenta", "Fecha", "Comprobante", "Descripción",
                  "Debe", "Haber", "Documento", "Dimensión valor", "Detalle", "Cotizacion",
                  "Productoid", "Producto", "Empresa", "Organizacion transaccion",
                  "Debe", "Haber", "Tipo de documento"]


def movimiento(codigo, centavos):
    """Un movimiento inventado del 30/09/2026 con ese Código y monto."""
    return Movimiento(date(2026, 9, 30), codigo, "", centavos)


def extracto_bbva(ruta, filas):
    """Guarda un extracto del BBVA inventado con esas filas (código, crédito, débito)."""
    libro = openpyxl.Workbook()
    libro.active.append(ENCABEZADO_BBVA)
    for codigo, credito, debito in filas:
        libro.active.append(["30-09-2026", "30-09-2026", "X", codigo, "", "OF", credito, debito, ""])
    libro.save(ruta)
    return ruta


def hojas(ruta):
    """Devuelve, por hoja del Excel, sus filas sin el encabezado."""
    libro = openpyxl.load_workbook(ruta)
    return {hoja.title: list(hoja.iter_rows(min_row=2, values_only=True)) for hoja in libro.worksheets}


def test_cada_fila_va_a_una_sola_lista_en_el_orden_del_original():
    reglas = {"gastos": ["236", "015"]}
    filas = [movimiento("319", 25000000), movimiento("236", -125000), movimiento("005", -15000000),
             movimiento("015", -95000), movimiento("319", 8000000)]
    sin_gastos, gastos = separar_gastos(filas, reglas)
    assert [m.codigo for m in sin_gastos] == ["319", "005", "319"]
    assert [m.codigo for m in gastos] == ["236", "015"]
    assert len(sin_gastos) + len(gastos) == len(filas)


def test_sin_gastos_todo_queda_en_movimientos():
    filas = [movimiento("319", 100), movimiento("005", -100)]
    assert separar_gastos(filas, {"gastos": ["236"]}) == (filas, [])


def test_el_extracto_sale_con_dos_hojas(tmp_path):
    # 236 y 015 son gasto en la lista real del BBVA; 319 y 005 no.
    ruta = extracto_bbva(tmp_path / "extracto.xlsx",
                         [("319", 2500, ""), ("236", "", -12.5), ("005", "", -1500), ("015", "", -9.5)])
    resumen = convertir_documento(ruta, "BBVA", False, "2026-09", tmp_path / "output")
    assert resumen["filas"] == 4 and resumen["movimientos"] == 2 and resumen["gastos"] == 2
    salida = hojas(resumen["ruta_salida"])
    assert list(salida) == ["Movimientos", "Gastos bancarios"]
    assert [(codigo, monto) for _, codigo, _, monto in salida["Movimientos"]] == [("319", 2500), ("005", -1500)]
    assert [(codigo, monto) for _, codigo, _, monto in salida["Gastos bancarios"]] == [("236", -12.5), ("015", -9.5)]


def test_las_dos_hojas_tienen_el_mismo_encabezado_y_el_codigo_como_texto(tmp_path):
    ruta = extracto_bbva(tmp_path / "extracto.xlsx", [("319", 100, ""), ("015", "", -1)])
    resumen = convertir_documento(ruta, "BBVA", False, "2026-09", tmp_path / "output")
    libro = openpyxl.load_workbook(resumen["ruta_salida"])
    for hoja in libro.worksheets:
        assert [celda.value for celda in hoja[1]] == ["Fecha", "Código", "Descripción", "Monto"]
    assert libro["Gastos bancarios"]["B2"].value == "015"


def test_un_mes_sin_gastos_tiene_la_hoja_de_gastos_vacia(tmp_path):
    ruta = extracto_bbva(tmp_path / "extracto.xlsx", [("319", 100, "")])
    resumen = convertir_documento(ruta, "BBVA", False, "2026-09", tmp_path / "output")
    assert resumen["gastos"] == 0
    assert hojas(resumen["ruta_salida"]) == {"Movimientos": [(datetime(2026, 9, 30), "319", None, 100)],
                                             "Gastos bancarios": []}


def test_el_libro_del_erp_sale_con_una_sola_hoja(tmp_path):
    libro = openpyxl.Workbook()
    libro.active.append(ENCABEZADO_ERP)
    libro.active.append(["2026-09", "1", "BANCO", datetime(2026, 9, 1), "F", "", 0, 1500, "PAGO - 1",
                         None, "", 1, None, None, "EMPRESA", "", 0, 1, "Orden de pago"])
    libro.save(tmp_path / "libro.xlsx")
    resumen = convertir_documento(tmp_path / "libro.xlsx", "BBVA", True, "2026-09", tmp_path / "output")
    assert "gastos" not in resumen
    assert list(hojas(resumen["ruta_salida"])) == ["Movimientos"]
